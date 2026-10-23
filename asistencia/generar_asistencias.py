import argparse
import calendar
import re
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


MESES = (
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
)
GRIS_FIN_DE_SEMANA = "7F7F7F"
BLANCO = "FFFFFF"


def argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera listas de asistencia quincenales desde una plantilla DOCX."
    )
    parser.add_argument(
        "--plantilla",
        type=Path,
        default=Path("ASISTENCIA(1).docx"),
        help="Documento DOCX que se usara como plantilla.",
    )
    parser.add_argument(
        "--salida",
        type=Path,
        default=Path("asistencias_2026"),
        help="Directorio donde se guardaran los documentos.",
    )
    parser.add_argument(
        "--inicio",
        help="Primer mes a generar en formato AAAA-MM. Por defecto usa el mes posterior a la plantilla.",
    )
    parser.add_argument("--nombre", help="Nombre que aparecera en las hojas de asistencia.")
    parser.add_argument("--puesto", help="Puesto que aparecera en las hojas de asistencia.")
    parser.add_argument("--adscripcion", help="Adscripcion que aparecera en las hojas de asistencia.")
    return parser.parse_args()


def periodo_de_plantilla(documento: Document) -> tuple[int, int]:
    encabezado = documento.paragraphs[0].text.upper()
    coincidencia_anio = re.search(r"A[ÑN]O\s*:\s*(\d{4})", encabezado)
    mes = next((numero for numero, nombre in enumerate(MESES, 1) if nombre in encabezado), None)
    if mes is None or coincidencia_anio is None:
        raise ValueError("No se pudo identificar MES y ANO en el primer parrafo de la plantilla.")
    return int(coincidencia_anio.group(1)), mes


def siguiente_mes(anio: int, mes: int) -> tuple[int, int]:
    return (anio + 1, 1) if mes == 12 else (anio, mes + 1)


def interpretar_inicio(valor: str | None, documento: Document) -> tuple[int, int]:
    if valor is None:
        return siguiente_mes(*periodo_de_plantilla(documento))
    coincidencia = re.fullmatch(r"(\d{4})-(0[1-9]|1[0-2])", valor)
    if coincidencia is None:
        raise ValueError("--inicio debe tener el formato AAAA-MM.")
    return int(coincidencia.group(1)), int(coincidencia.group(2))


def reemplazar_periodo(documento: Document, anio: int, mes: int) -> None:
    parrafo = documento.paragraphs[0]
    cambio_mes = False
    cambio_anio = False
    for run in parrafo.runs:
        if any(nombre in run.text.upper() for nombre in MESES):
            espacios = run.text[: len(run.text) - len(run.text.lstrip())]
            run.text = f"{espacios}{MESES[mes - 1]}"
            cambio_mes = True
        if re.search(r"\b\d{4}\b", run.text):
            run.text = re.sub(r"\b\d{4}\b", str(anio), run.text)
            cambio_anio = True
    if not cambio_mes or not cambio_anio:
        raise ValueError("No se pudieron reemplazar MES y ANO sin alterar el formato.")


def establecer_texto_celda(celda, texto: str) -> None:
    parrafo = celda.paragraphs[0]
    if parrafo.runs:
        parrafo.runs[0].text = texto
        for run in parrafo.runs[1:]:
            run.text = ""
    else:
        parrafo.add_run(texto)


def establecer_sombreado(celda, color: str) -> None:
    propiedades = celda._tc.get_or_add_tcPr()
    for sombreado_anterior in propiedades.findall(qn("w:shd")):
        propiedades.remove(sombreado_anterior)
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:color"), "auto")
    sombreado.set(qn("w:fill"), color)
    propiedades.append(sombreado)


def ajustar_filas(tabla, cantidad_dias: int) -> None:
    filas_datos = len(tabla.rows) - 1
    while filas_datos < cantidad_dias:
        tabla._tbl.append(deepcopy(tabla.rows[-1]._tr))
        filas_datos += 1
    while filas_datos > cantidad_dias:
        tabla._tbl.remove(tabla.rows[-1]._tr)
        filas_datos -= 1


def llenar_quincena(documento: Document, anio: int, mes: int, quincena: int) -> None:
    if len(documento.tables) != 1 or len(documento.tables[0].columns) != 6:
        raise ValueError("La plantilla debe contener una tabla de asistencia de seis columnas.")
    ultimo_dia = calendar.monthrange(anio, mes)[1]
    dias = list(range(1, 16)) if quincena == 1 else list(range(16, ultimo_dia + 1))
    tabla = documento.tables[0]
    ajustar_filas(tabla, len(dias))
    for fila, dia in zip(tabla.rows[1:], dias):
        establecer_texto_celda(fila.cells[0], str(dia))
        es_fin_de_semana = date(anio, mes, dia).weekday() >= 5
        color = GRIS_FIN_DE_SEMANA if es_fin_de_semana else BLANCO
        establecer_sombreado(fila.cells[0], color)
        for celda in fila.cells[1:]:
            establecer_texto_celda(celda, "")
            establecer_sombreado(celda, color)


def reemplazar_datos(
    documento: Document, nombre: str | None, puesto: str | None, adscripcion: str | None,
) -> None:
    for etiqueta, valor in (("NOMBRE", nombre), ("PUESTO", puesto), ("ADSCRIPCIÓN", adscripcion)):
        if valor is None:
            continue
        if not valor.strip() or any(caracter in valor for caracter in "\r\n\t"):
            raise ValueError(f"{etiqueta} debe contener texto en una sola linea.")
        encontrado = False
        for parrafo in documento.paragraphs:
            coincidencia = re.search(rf"{etiqueta}:\s*([^\t\r\n]*\S)", parrafo.text)
            if coincidencia is None:
                continue
            inicio, fin = coincidencia.span(1)
            posicion = 0
            insertado = False
            for run in parrafo.runs:
                texto = run.text
                limite = posicion + len(texto)
                if posicion < fin and limite > inicio:
                    prefijo = texto[:max(0, inicio - posicion)]
                    sufijo = texto[max(0, fin - posicion):]
                    run.text = prefijo + (valor.strip() if not insertado else "") + sufijo
                    insertado = True
                posicion = limite
            encontrado = True
        if not encontrado:
            raise ValueError(f"No se encontro el campo {etiqueta} en la plantilla.")


def crear_documento(
    plantilla: Path, anio: int, mes: int, quincena: int,
    nombre: str | None = None, puesto: str | None = None, adscripcion: str | None = None,
) -> Document:
    documento = Document(plantilla)
    reemplazar_datos(documento, nombre, puesto, adscripcion)
    reemplazar_periodo(documento, anio, mes)
    llenar_quincena(documento, anio, mes, quincena)
    return documento


def generar(
    plantilla: Path, salida: Path, inicio: tuple[int, int],
    nombre: str | None = None, puesto: str | None = None, adscripcion: str | None = None,
) -> list[Path]:
    anio, mes_inicial = inicio
    salida.mkdir(parents=True, exist_ok=True)
    archivos = []
    for mes in range(mes_inicial, 13):
        for quincena in (1, 2):
            documento = crear_documento(plantilla, anio, mes, quincena, nombre, puesto, adscripcion)
            destino = salida / f"ASISTENCIA_{anio}_{mes:02d}_Q{quincena}.docx"
            documento.save(destino)
            archivos.append(destino)
    return archivos


def main() -> None:
    args = argumentos()
    if not args.plantilla.is_file():
        raise FileNotFoundError(f"No existe la plantilla: {args.plantilla}")
    documento = Document(args.plantilla)
    inicio = interpretar_inicio(args.inicio, documento)
    archivos = generar(args.plantilla, args.salida, inicio, args.nombre, args.puesto, args.adscripcion)
    print(f"Se generaron {len(archivos)} documentos en {args.salida.resolve()}")
    for archivo in archivos:
        print(archivo)


if __name__ == "__main__":
    main()